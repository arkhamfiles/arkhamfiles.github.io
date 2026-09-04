#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""아컴 마도서 번역대조표(xlsx) → chapters.html 챕터 2 패널 재생성 스크립트.

사용법:  python grimoire_generator.py <번역대조표.xlsx> <chapters.html>
 - 챕터 2의 규칙/정오표/FAQ/선택 규칙/개정 재수록 패널(본문+목차)을 시트 내용으로 교체합니다.
 - 시트가 갱신되면 다시 실행하기만 하면 됩니다. (chapters.html은 제자리에서 수정)
"""
import os, sys, re, html, unicodedata
import openpyxl

ICON = {
 'free':('icon-free','자유 격발'), 'action':('icon-action','행동 격발'), 'reaction':('icon-reaction','반응 격발'),
 'combat':('icon-combat','힘'), 'intellect':('icon-intellect','지식'), 'willpower':('icon-willpower','의지'),
 'agility':('icon-agility','민첩'), 'wild':('icon-wild','만능'), 'skull':('icon-skull','해골'),
 'cultist':('icon-cultist','추종자'), 'tablet':('icon-tablet','석판'), 'elder_thing':('icon-elder_thing','고대의 것'),
 'elder_sign':('icon-elder_sign','고대 표식'), 'auto_fail':('icon-auto_fail','자동 실패'),
 'bless':('icon-bless','축복'), 'curse':('icon-curse','저주'), 'unique':('icon-unique','고유'),
 'per_investigator':('icon-per_investigator','조사자당'), 'inv':('icon-per_investigator','조사자당'),
 'guardian':('icon-guardian','수호자'), 'seeker':('icon-seeker','탐구자'), 'rogue':('icon-rogue','무법자'),
 'mystic':('icon-mystic','신비주의자'), 'survivor':('icon-survivor','생존자'),
 'core26':('sym2-core26','기본판(2026)'), 'core2026':('sym2-core26','기본판(2026)'), 'fight':('icon-combat','힘'),
 'tommy':('sym2-tommy','토미 멀둔'), 'carolyn':('sym2-carolyn','캐롤린 펀'), 'andre':('sym2-andre','앙드레 파텔'),
 'marie':('sym2-marie','마리 램부'), 'maire':('sym2-marie','마리 램부'), 'miguel':('sym2-miguel','미겔 데 라 크루스'), 'codex':('sym2-codex','서고'),
 'weakness':('sym2-weakness','기본 약점'), 'basic_weakness':('sym2-weakness','기본 약점'),
}
DROP_MARKERS = set()

def fmt(text):
    if text is None: return ''
    s = html.escape(str(text).strip(), quote=False)
    def rep(m):
        k = m.group(1)
        if k in DROP_MARKERS: return ''
        if k in ICON:
            cls, t = ICON[k]
            return f'<span title="{t}" class="{cls}"></span>'
        return m.group(0)
    s = re.sub(r'\[([a-z_0-9]+)\]', rep, s)
    s = re.sub(r'\*\*\*(.+?)\*\*\*', r'<b><i>\1</i></b>', s)
    s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s)
    s = re.sub(r'\*(.+?)\*', r'<i>\1</i>', s)
    s = re.sub(r'^(알아두기|참고|예시|주의):', r'<b>\1:</b>', s)
    return s

def slug(s):
    s = unicodedata.normalize('NFKD', str(s))
    s = re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_').lower()
    return s[:48] or 'x'

def cell_text(v):
    """셀 값 → 문자열. 리치 텍스트(부분 볼드/이탤릭)는 **·*·*** 마커로 변환"""
    try:
        from openpyxl.cell.rich_text import CellRichText
    except Exception:
        CellRichText = ()
    if isinstance(v, CellRichText):
        out=[]
        for run in v:
            t = str(run)
            f = getattr(run, 'font', None)
            b = bool(getattr(f, 'b', False)); i = bool(getattr(f, 'i', False))
            if t.strip()=='' or (not b and not i): out.append(t); continue
            lead = t[:len(t)-len(t.lstrip())]; trail = t[len(t.rstrip()):]
            core = t.strip()
            mk = '***' if (b and i) else ('**' if b else '*')
            out.append(lead+mk+core+mk+trail)
        return ''.join(out)
    return v

def parse(xlsx):
    wb = openpyxl.load_workbook(xlsx, rich_text=True)
    ws = wb['마도서 대조표']
    rows = []
    cur_sec = None
    for cr in ws.iter_rows(min_row=2):
        vals = [cell_text(c.value) for c in cr[:6]]
        # 번역(한글) 셀: 셀 전체에 걸린 볼드/이탤릭(런이 아닌 셀 서식)도 마커로 반영
        if len(cr) > 5 and isinstance(vals[5], str) and vals[5].strip():
            f = cr[5].font
            _plain = vals[5].strip()
            structural = _plain.startswith(('Æ','= ','> ','참조:'))
            if f is not None and '**' not in vals[5] and not structural:
                b = bool(f.b); it = bool(f.i)
                if b or it:
                    mk = '***' if (b and it) else ('**' if b else '*')
                    vals[5] = mk + vals[5].strip() + mk
        page, sec, et, kt, en, ko = (tuple(vals) + (None,)*6)[:6]
        norm = lambda v: str(v).strip() if v not in (None, 'None') and str(v).strip() not in ('','None') else None
        sec, et, kt, en, ko = norm(sec), norm(et), norm(kt), norm(en), norm(ko)
        explicit = bool(sec)
        if sec: cur_sec = sec
        rows.append(dict(sec=cur_sec, et=et, kt=kt, en=en, ko=ko, sec_explicit=explicit))
    return rows

def sec_key(sec):
    if not sec: return None
    s = sec
    for pref,key in [('Cover','cover'),('Overview','overview'),('I. Glossary','glossary'),
        ('II. Additional','fundamentals'),('III. Timing','timing'),('IV. Skill','skilltest'),
        ('V. Initiation','initiation'),('VI. Scenario','setup'),('VII. Card','anatomy'),
        ('VIII. Campaign','campaign'),('IX. Deck','deckcustom'),('X. Notes','errata'),
        ('XI. Freq','faq'),('XII. Optional','optional'),('XIII. Modified','modreprints'),
        ('XIV. Icon','iconref'),('XV. Encounter','encicon'),
        ('Encounter Set Variation','encvar'),('Quick Reference','quickref'),('Index','index')]:
        if s.startswith(pref): return key
    return None

SKIP_SECTIONS = {'cover', 'index', None}
SKIP_ET = {'(조우 세트 대체표 라벨 — 페이지 상단 파편)', 'Contents', 'Table of Contents'}
SKIP_KT = {'목차'}
def classify(row):
    ko, en, kt = row['ko'], row['en'], row['kt']
    if not ko: return None
    if kt == '>' or ko.startswith('> '):
        return ('cont', ko[2:].strip() if ko.startswith('> ') else ko)
    _r=ref_rest(ko) if 'REF_RE' in globals() else None
    if _r is None and ko.startswith('참조:'): _r=ko[3:].strip()
    if _r is not None:
        return ('ref', _r)
    if kt == '=' or ko.startswith('= '):
        return ('li2', ko[2:].strip() if ko.startswith('= ') else ko)
    if (kt == 'Æ') or (en and en.startswith('Æ ')) or ko.startswith('Æ '):
        return ('li', ko[2:].strip() if ko.startswith('Æ ') else ko)
    if re.fullmatch(r'\d{1,2}\.', ko):
        return ('num', ko)
    if re.match(r'^.+ → .+$', ko) and len(ko) < 80:
        return ('arrow', ko)
    return ('p', ko)

CARDTYPE_SUBS={'Act Cards','Agenda Cards','Asset Cards','Enemy Cards','Event Cards','Investigator Cards','Location Cards','Scenario Reference Cards','Skill Cards','Story Cards','Treachery Cards'}

def build_entries(rows):
    sections = {}; order = []; cur = None; override = None
    for row in rows:
        if row.get('sec_explicit'): override = None
        k = sec_key(row['sec'])
        if override: k = override
        if row.get('et'):
            k2 = sec_key(row['et'])
            if k2 and k2 != k:
                k = k2; override = k2
        if k in SKIP_SECTIONS: cur=None; continue
        if row['et'] and (row['et'] in SKIP_ET or (row['kt'] or '') in SKIP_KT): cur=None; continue
        is_hdr = bool(row['et'] and re.match(r'^(X{0,3}(IX|IV|V?I{0,3})\.\s|Overview$|Quick Reference / )', row['et']))
        if k not in sections:
            sections[k] = dict(key=k, title=None, intro=[], entries=[]); order.append(k)
        S = sections[k]
        S.setdefault('rows',[]).append(dict(row, is_hdr=bool(is_hdr)))
        if is_hdr:
            if row['kt']:
                S['title'] = re.sub(r'^[IVX]+\.\s*', '', row['kt'].split(' / ')[-1] if ' / ' in row['kt'] else row['kt'])
            if row['ko']:
                p=classify(row)
                if p: S['intro'].append(p)
            cur = ('intro', S); continue
        if row['et']:
            # 페이지 넘김으로 같은 표제어가 다시 나온 경우: 이전 엔트리에 병합
            if cur and cur[0]=='entry' and cur[1]['et']==row['et']:
                if row['ko'] and row['ko'] != (row['kt'] or ''):
                    p=classify(row)
                    if p: cur[1]['parts'].append(p)
                continue
            # '카드 종류' 항목에 딸린 카드 종류별 요약: 서브카드로 중첩
            if cur and cur[0]=='entry' and row['et'] in CARDTYPE_SUBS and                (cur[1]['et']=='Cardtypes' or cur[1].get('parent')=='Cardtypes'):
                sub=dict(kt=row['kt'] or row['et'], parts=[], parent='Cardtypes')
                host=cur[1] if cur[1]['et']=='Cardtypes' else cur[2]
                host.setdefault('subs',[]).append(sub)
                cur=('entry', dict(et=row['et'], parts=sub['parts'], parent='Cardtypes'), host)
                if row['ko'] and row['ko'] != (row['kt'] or ''):
                    p=classify(row)
                    if p: sub['parts'].append(p)
                continue
            e = dict(key='c2-'+k[:4]+'-'+slug(row['et']), et=row['et'], kt=row['kt'] or row['et'], parts=[])
            S['entries'].append(e); cur = ('entry', e)
            if row['ko'] and row['ko'] != (row['kt'] or ''):
                p=classify(row)
                if p: e['parts'].append(p)
            continue
        if cur is None: continue
        part = classify(row)
        if part is None: continue
        (S['intro'] if cur[0]=='intro' else cur[1]['parts']).append(part) if cur[0]=='intro' else cur[1]['parts'].append(part)
    return sections, order

TITLE_MAP = {}   # 현재 렌더링 중인 패널의 '제목 -> #앵커' 맵
CURRENT_ENTRY = None   # 렌더링 중인 엔트리 id (참조 예외 판정용)

# 참조 링크 예외조항: (엔트리 id, 참조 이름) -> 강제 대상 앵커
REF_EXCEPTIONS = {
    # 용어 해설의 '캠페인 플레이' 스텁은 자가 인용이 아니라 VIII. 캠페인 플레이 섹션으로
    ('c2-glos-campaign_play', '캠페인 플레이'): '#c2sec-campaign_play',
}

def strip_markers(t):
    return re.sub(r'\[[a-z_0-9]+\]\s*', '', t or '').strip()

def norm_quotes(t):
    return re.sub(r'[“”‘’"\']', '', t)

def set_title_map(secs):
    global TITLE_MAP
    TITLE_MAP = {}
    prim=[]
    for sid, sdata in secs:
        t = strip_markers(sdata.get('title') or '')
        if t: prim.append((t, '#'+sid))
        for e in sdata.get('entries', []):
            prim.append((strip_markers(e['kt']), '#'+e['key']))
    # 1차: 원제목 우선 등록
    for t, tgt in prim: TITLE_MAP.setdefault(t, tgt)
    # 2차: 변형(괄호 제거·따옴표 제거·결합 제목의 부분 명칭) 등록
    for t, tgt in prim:
        variants=set()
        base=re.sub(r'\s*\([^)]*\)$','',t)
        variants.update([base, norm_quotes(t), norm_quotes(base)])
        for part in re.split(r',\s*', t):           # '회피, 회피 행동'
            p=part.strip()
            if len(p)>=2: variants.add(p)
        m=re.match(r'^(.{2,}?[덱벌섯록])[과와]\s(.{2,})$', t)  # '주요목적 덱과 주요사건 덱'
        if m: variants.update([m.group(1), m.group(2)])
        for v in variants:
            if v and v!=t: TITLE_MAP.setdefault(v, tgt)

def lookup_ref(n):
    cands=[n, norm_quotes(n), re.sub(r'\s*\([^)]*\)$','',n),
           norm_quotes(re.sub(r'\s*\([^)]*\)$','',n)),
           re.sub(r'^(\d+\.\d+):\s*', r'\1 ', n)]
    for c in cands:
        c=c.strip()
        if c in TITLE_MAP: return TITLE_MAP[c]
    return None

def linkify_refs(text):
    out=[]
    for name in re.split(r',\s*', text):
        n=name.strip()
        if not n: continue
        tgt = REF_EXCEPTIONS.get((CURRENT_ENTRY, strip_markers(n))) or lookup_ref(strip_markers(n))
        out.append('<a href="'+tgt+'">'+fmt(n)+'</a>' if tgt else fmt(n))
    return ', '.join(out)


REF_RE = re.compile(r"^\*{0,3}참조\s*:\s*\*{0,3}\s*")

def ref_rest(text):
    m = REF_RE.match(text)
    return REF_RE.sub("", text, count=1).strip() if m else None

def cont_html(text, cls):
    """이어쓰기(>) 내용 렌더링: 참조 행이면 참조 서식+링크 유지"""
    rest=ref_rest(text)
    if rest is not None:
        inner='<span>참조:</span> '+linkify_refs(rest)
        return '<div class="block-refs'+((' '+cls) if cls else '')+'">'+inner+'</div>'
    return '<p class="'+(cls or 'cont-plain')+'">'+fmt(text)+'</p>'

def render_parts(parts):
    out=[]; i=0
    while i < len(parts):
        kind, text = parts[i]
        if kind=='li2':
            items=[]
            while i<len(parts) and parts[i][0] in ('li2','cont'):
                k2,t2=parts[i]; i+=1
                if k2=='li2': items.append([t2,[]])
                elif items: items[-1][1].append(t2)
            lis=''.join('<li>'+fmt(t)+''.join(cont_html(c,'licont') for c in cs)+'</li>' for t,cs in items)
            out.append('<ul class="subnum">'+lis+'</ul>'); continue
        if kind in ('li','li2'):
            groups=[]; last=None
            while i<len(parts) and parts[i][0] in ('li','li2','cont'):
                k2,t2=parts[i]; i+=1
                if k2=='li' or (k2=='li2' and not groups):
                    groups.append({'t':t2,'conts':[],'subs':[]}); last=groups[-1]
                elif k2=='li2':
                    groups[-1]['subs'].append({'t':t2,'conts':[]}); last=groups[-1]['subs'][-1]
                elif last is not None:
                    last['conts'].append(t2)
            def li_html(g):
                inner=fmt(g['t'])+inline_imgs_html(g['t'])+''.join(cont_html(c,'licont') for c in g['conts'])
                if g.get('subs'):
                    inner+='<ul>'+''.join(li_html(sg) for sg in g['subs'])+'</ul>'
                return '<li>'+inner+'</li>'
            out.append('<ul>'+''.join(li_html(g) for g in groups)+'</ul>'); continue
        if kind=='num':
            items=[]
            while i<len(parts) and parts[i][0]=='num':
                n=parts[i][1]; i+=1
                body = fmt(parts[i][1]) if i<len(parts) and parts[i][0]=='p' else ''
                if body: i+=1
                items.append('<li value="'+n.rstrip('.')+'">'+body+'</li>')
            out.append('<ol>'+''.join(items)+'</ol>'); continue
        if kind=='arrow':
            trs=[]
            while i<len(parts) and parts[i][0]=='arrow':
                a,b=[fmt(x.strip()) for x in parts[i][1].split('→',1)]; i+=1
                trs.append('<tr><td>'+a+'</td><td>'+b+'</td></tr>')
            out.append('<table><tr><th>기존 조우 세트</th><th>대체 조우 세트</th></tr>'+''.join(trs)+'</table>'); continue
        if kind=='q':
            out.append('<p class="question">'+fmt(text)+'</p>'); i+=1; continue
        if kind=='a':
            out.append('<p class="answer">'+fmt(text)+'</p>'); i+=1; continue
        if kind=='h4':
            sid,t=text
            out.append('<h4 class="fw-step" id="'+sid+'">'+fmt(t)+'</h4>'); i+=1; continue
        if kind=='chart':
            key,nodes=text
            hn=[]
            for idx,nd in enumerate(nodes):
                if idx: hn.append('<div class="fc-arr'+(' fc-arr-r' if nd['t']=='w' else '')+'"></div>')
                cls='fc-node'+(' fc-window' if nd['t']=='w' else '')+(' fc-end' if nd['t']=='e' else '')
                hn.append('<div class="'+cls+'">'+fmt(nd['text'])+'</div>')
            loops=''.join('<div class="fc-loop" data-from="%d" data-to="%d"></div>'%(a,b) for a,b in CHART_LOOPS.get(key,[]))
            out.append('<div class="fc">'+''.join(hn)+loops+'</div>'); i+=1; continue
        if kind=='ref':
            out.append('<div class="block-refs"><span>참조:</span> '+linkify_refs(text)+'</div>'); i+=1; continue
        if kind=='cont':
            j=i-1
            while j>=0 and parts[j][0]=='cont': j-=1
            prev_num = j>=0 and parts[j][0]=='p' and re.match(r'^\d{1,2}\.\s', parts[j][1])
            out.append(cont_html(text, 'cont-num' if prev_num else 'cont-plain')); i+=1; continue
        cls=' class="numbered"' if re.match(r'^\d{1,2}\.\s', text) else ''
        out.append('<p'+cls+'>'+fmt(text)+'</p>'); i+=1
        out.append(inline_imgs_html(text))
        # 연속 cont가 번호 문단을 따라오도록 표시 유지
    return ''.join(out)

def entry_html(e):
    global CURRENT_ENTRY
    CURRENT_ENTRY = e.get('key')
    if e.get('stopbox'):
        lines=' <br /> '.join(fmt(t) for k,t in e['parts'] if k in ('p','cont'))
        return ('<article class="rules-block plain entry" id="'+e["key"]+'">'
                '<div class="block-body"><p class="stopMsgBox"><span>'+fmt(e["kt"])+'</span> <br /> '
                +lines+'</p></div></article>')
    # 제목과 동일한 도입 문단("...:" 형태) 중복 제거
    if e['parts'] and e['parts'][0][0]=='p':
        t=e['parts'][0][1].rstrip(':').strip()
        if t==e['kt'].strip(): e['parts']=e['parts'][1:]
    subs=''.join('<article class="rules-block subcard"><h3 class="block-title">'+fmt(s2["kt"])+'</h3>'
                 '<div class="block-body">'+render_parts(s2["parts"])+'</div></article>'
                 for s2 in e.get('subs',[]))
    tcls='block-title fc-title' if e.get('center_title') else 'block-title'
    title = '' if e.get('notitle') else '<h3 class="'+tcls+'">'+fmt(e["kt"])+'</h3>'
    cls='rules-block entry'+(' plain' if e.get('plain') else '')
    row = e['key'] in ENTRY_IMAGES_ROW
    imgs=''
    for it in ENTRY_IMAGES.get(e['key'],[]):
        u,xcls = it if isinstance(it,tuple) else (it,'')   # (경로, 추가클래스) 또는 경로
        imgs+=('<figure class="rb-fig'+(' '+xcls if xcls else '')+'"'+(' style="flex-grow:%.4f"'%img_aspect(u) if row else '')+
               '><img loading="lazy" src="'+u+'" alt="'+html.escape(e["kt"],quote=True)+'"></figure>')
    if imgs and row: imgs='<div class="rb-row">'+imgs+'</div>'
    body=render_parts(e["parts"])+subs
    if e['key'] in IMAGE_REPLACES_BODY and imgs: body=''
    return ('<article class="'+cls+'" id="'+e["key"]+'">'+title+
            '<div class="block-body">'+body+imgs+'</div></article>')

def section_html(sec_id, title, intro, entries):
    h = '<section class="ch2-section" id="'+sec_id+'"><h2 class="subsec-title">'+fmt(title)+'</h2>'
    if intro:
        h += '<div class="block-body sec-intro">'+render_parts(intro)+'</div>'
    h += '</section>\n'
    h += '\n'.join(entry_html(e) for e in entries)
    return h

def toc_section(sec_id, title, entries):
    def entry_items(e):
        h=''
        if not e.get('notoc'):
            h+='<a class="idx-item" href="#'+e["key"]+'" data-title="'+html.escape((e["kt"]+' '+e["et"]).lower(),quote=True)+'">'+fmt(e["kt"])+'</a>'
        subcls='idx-item' if e.get('subitems_base') else 'idx-item idx-sub'
        for sid,st in e.get('subitems',[]):
            h+='<a class="'+subcls+'" href="#'+sid+'" data-title="'+html.escape(st.lower(),quote=True)+'">'+fmt(st)+'</a>'
        return h
    items=''.join(entry_items(e) for e in entries)
    return ('<div class="idx-section" data-section="'+sec_id+'">'
            '<div class="idx-section-title"><a class="sec-jump" href="#'+sec_id+'">'+fmt(title)+'</a></div>'+items+'</div>')

def sort_kr(entries):
    return sorted(entries, key=lambda e: re.sub(r'^["“‘\'『〈<\(]+','',e['kt']))
def split_faq(sections):
    """XI. FAQ: 표제어 없는 질문/답변 연속 행을 Q&A 쌍 카드로 분할"""
    S = sections.get('faq')
    if not S: return
    parts = list(S.get('intro', []))
    # 질문 행이 나오기 전까지는 섹션 도입문으로 유지
    intro, qa, seen_q = [], [], False
    for kind, text in parts:
        is_q = kind=='p' and re.sub(r'\*+','',text).rstrip().endswith('?')
        if is_q: seen_q = True
        (qa if seen_q else intro).append((kind, text))
    S['intro'] = intro
    entries=[]; cur=None; n=0
    for kind, text in qa:
        is_q = kind=='p' and re.sub(r'\*+','',text).rstrip().endswith('?')
        if is_q:
            n+=1
            disp=re.sub(r'\[[a-z_0-9]+\]\s*','',text)   # 목차 표시용: 아이콘 마커 제거 후 절단
            cur=dict(key='c2-faq-q%02d'%n, et='FAQ Q%d'%n,
                     kt=(disp[:42]+'…') if len(disp)>42 else disp,
                     parts=[('q',text)], notitle=True)
            entries.append(cur)
        elif cur is not None:
            cur['parts'].append(('a',text) if kind=='p' else (kind,text))
    S['entries']=entries

def split_cardtype_subs(sections):
    """'카드 종류' 항목: 'OO 카드' 단독 행을 하위 카드 제목으로 인식해 서브카드로 분할"""
    g = sections.get('glossary')
    if not g: return
    e = next((x for x in g['entries'] if x.get('et')=='Cardtypes'), None)
    if not e: return
    intro=[]; subs=[]; cur=None
    for kind, text in e['parts']:
        plain = re.sub(r'\*+','',text).strip()
        if kind=='p' and len(plain)<=14 and plain.endswith('카드') and not any(c in plain for c in '.,:()'):
            cur=dict(kt=plain, parts=[]); subs.append(cur)
        elif cur is None: intro.append((kind,text))
        else: cur['parts'].append((kind,text))
    if subs:
        e['parts']=intro; e['subs']=subs

CHART_LOOPS = {  # 차트 첫 단계 번호 -> [(from노드, to노드)] 회귀 화살표
    '2': [(4,3),(5,2)],   # 조사 단계: 2.2.1→직전 기회, 2.2.2→2.2
    '3': [(3,2)],         # 적 단계: 3.3→직전 기회
}

def _is_window(ko): return bool(ko) and re.sub(r'^\*+','',ko).strip().startswith('플레이어의') and '기회' in ko
def _is_chart_step(row):
    en=(row.get('en') or ''); ko=(row.get('ko') or '')
    return en.startswith('Step') or 'Player Window' in en or bool(re.match(r'^\*{0,3}(과정\s|\d단계:)', ko))
def _is_phase_title(ko): return bool(re.match(r'^\*{0,3}[IVX]+\.\s', ko or ''))

def collect_chart(rows, i):
    """rows[i]가 차트 헤더일 때 (노드목록, 다음 인덱스) 반환"""
    nodes=[]; j=i+1
    while j < len(rows):
        r=rows[j]; ko=r.get('ko'); en=(r.get('en') or '')
        if not ko and not en: j+=1; continue
        if _is_window(ko):
            nodes.append(dict(t='w', text='플레이어의 [free] 기회')); j+=1; continue
        if _is_chart_step(r) and ko:
            nodes.append(dict(t='s', text=ko)); j+=1; continue
        if ko and re.sub(r'\*+','',ko).rstrip().endswith('이어집니다.'):
            nodes.append(dict(t='e', text=ko)); j+=1; continue
        break
    return nodes, j

def chart_entry(header_row, nodes, seckey):
    ko=header_row.get('ko') or '차트'
    m=re.search(r'(과정\s)?(\d)', ''.join(n['text'] for n in nodes[:1]))
    key=m.group(2) if m else ''
    disp=re.sub(r'^\*{0,3}[IVX]+\.\s*','',ko)          # 로마 숫자 제거
    return dict(key='c2-'+seckey[:4]+'-chart_'+slug(header_row.get('en') or ko),
                et=(header_row.get('en') or ko)+' (chart)', kt=disp+' 차트',
                parts=[('chart',(key,nodes))], notoc=True, center_title=True)

def restructure_timing(sections):
    S=sections.get('timing')
    if not S or not S.get('rows'): return
    rows=S['rows']
    # 도입부(섹션 헤더~첫 ET 전)는 기존 intro 유지, 첫 ET(잠깐!!)부터 재구성
    start=next((i for i,r in enumerate(rows) if r.get('et') and not r.get('is_hdr')), None)
    if start is None: return
    entries=[]; cur=None; i=start
    while i < len(rows):
        r=rows[i]; ko=r.get('ko'); en=(r.get('en') or '')
        if not ko and not en: i+=1; continue
        if _is_phase_title(ko):
            nxt=next((rows[j] for j in range(i+1,len(rows)) if rows[j].get('ko') or rows[j].get('en')), None)
            if nxt is not None and _is_chart_step(nxt):
                nodes,i2=collect_chart(rows,i)
                entries.append(chart_entry(r,nodes,'timing')); cur=None; i=i2; continue
            cur=dict(key='c2-timi-'+slug(en or ko), et=en or ko, kt=ko, parts=[], subitems=[])
            entries.append(cur); i+=1; continue
        if r.get('et'):  # 잠깐!! 등 일반 표제어
            cur=dict(key='c2-timi-'+slug(r['et']), et=r['et'], kt=r.get('kt') or r['et'], parts=[])
            if (cur['kt'] or '').startswith('잠깐'): cur['stopbox']=True; cur['notoc']=True
            entries.append(cur)
            p=classify(r)
            if p and (r.get('ko') or '') != (r.get('kt') or ''): cur['parts'].append(p)
            i+=1; continue
        if re.match(r'^\*{0,3}\d\.\d(\.\d)?\s', ko or '') and not en.startswith('Step'):
            sid='c2-timi-'+slug(en or ko)
            if cur is not None and 'subitems' in cur:   # 단계 카드 안의 과정 소제목(목차에는 미표시)
                cur['parts'].append(('h4',(sid,ko)))
            else:
                cur=dict(key=sid, et=en or ko, kt=ko, parts=[])
                entries.append(cur)
            i+=1; continue
        p=classify(r)
        if p and cur is not None: cur['parts'].append(p)
        i+=1
    S['entries']=entries

def restructure_skilltest(sections):
    S=sections.get('skilltest')
    if not S or not S.get('rows'): return
    rows=S['rows']; entries=[]; cur=None; i=0
    while i < len(rows):
        r=rows[i]; ko=r.get('ko'); en=(r.get('en') or '')
        if not ko and not en: i+=1; continue
        if en=='Skill Test Timing' or (ko=='능력 테스트 순서와 시점' and not r.get('is_hdr')):
            nodes,i2=collect_chart(rows,i)
            if nodes:
                entries.append(chart_entry(r,nodes,'skilltest')); cur=None; i=i2; continue
        if en.startswith('ST.'):
            sid='c2-skil-'+slug(en)
            if not entries or 'subitems' not in entries[0]:
                host=dict(key='c2-skil-details', et='Skill Test Steps', kt='', parts=[], subitems=[],
                          notitle=True, notoc=True, subitems_base=True)
                entries.insert(0,host)
            host=entries[0]
            host['parts'].append(('h4',(sid,ko)))
            cur=host; i+=1; continue
        p=classify(r)
        if p and cur is not None: cur['parts'].append(p)
        i+=1
    S['entries']=entries
    S['intro']=[]

def tune_overview(sections):
    ov=sections.get('overview')
    if not ov: return
    ents=ov.get('entries',[])
    flavor=next((e for e in ents if (e.get('et') or '').startswith('Dark Secrets')), None)
    usage=next((e for e in ents if e.get('et')=='Using This Book'), None)
    stop=next((e for e in ents if (e.get('kt') or '').startswith('잠깐')), None)
    if usage:
        ov['title']=usage['kt']                       # 섹션 제목으로 격상
        usage['plain']=True; usage['notitle']=True; usage['notoc']=True
    if flavor:
        flavor['plain']=True; flavor['notoc']=True
        ov['pre']=[flavor]                            # 섹션 제목 위(인용문)
        ents.remove(flavor)
    if stop: stop['notoc']=True; stop['stopbox']=True

ANATOMY_IDS={'시나리오 카드 해설':'scenario_card_anatomy_key','플레이어 카드 해설':'player_card_anatomy_key'}

# 항목에 삽입할 규칙서 도해 이미지 (entry id -> [이미지 경로])
ENTRY_IMAGES = {
    'c2-anat-scenario_card_anatomy_key': ['images/rulebook/anatomy_scenario_1.webp','images/rulebook/anatomy_scenario_2.webp'],
    'c2-anat-player_card_anatomy_key':   [('images/rulebook/anatomy_player_1.webp','w75'),'images/rulebook/anatomy_player_2.webp'],  # w75: 75% 너비·가운데
}
IMAGE_REPLACES_BODY = set()
# 이미지를 가로로 나란히 배치할 항목
ENTRY_IMAGES_ROW = {'c2-anat-scenario_card_anatomy_key'}

def img_aspect(rel):
    """가로/세로 비율. 나란히 배치 시 flex-grow 로 써서 두 이미지의 높이를 맞춘다."""
    try:
        from PIL import Image
        w,h = Image.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), rel)).size
        return w/h
    except Exception:
        return 1.0
DROP_ENTRIES = {'Symbols and Icons'}   # 웹판에서 통째로 제외하는 항목

# 특정 문단 바로 아래 삽입하는 이미지: (엔트리 id, 문단에 포함된 문구) -> [(src, 추가클래스)]
INLINE_IMAGES = {
    ('c2-deck-random_basic_weaknesses', '다음과 같은 기호가 표시되어 있습니다'):
        [('images/rulebook/weakness_example.webp', 'plainfig')],
    ('c2-glos-weakness', '아래의 기호가 그려져 있고'):
        [('images/rulebook/basic_weakness_icon.webp', 'iconfig')],
}

def inline_imgs_html(text):
    h=''
    for (ek, frag), imgs in INLINE_IMAGES.items():
        if ek == CURRENT_ENTRY and frag in text:
            for src, xcls in imgs:
                h+='<figure class="rb-fig '+xcls+'"><img loading="lazy" src="'+src+'" alt=""></figure>'
    return h

def tune_anatomy(sections):
    """VII. 카드 해설: 도입부에 흡수된 'OO 해설' 단독 행을 항목으로 승격"""
    an=sections.get('anatomy')
    if not an: return
    intro=list(an.get('intro',[]))
    keep=[]; cur=None; new_entries=[]
    for k,t in intro:
        plain=re.sub(r'\*+','',t).strip() if isinstance(t,str) else ''
        if k=='p' and len(plain)<=14 and plain.endswith('해설') and '.' not in plain:
            cur=dict(key='c2-anat-'+ANATOMY_IDS.get(plain, slug(plain) or 'anat%d'%(len(new_entries)+1)),
                     et=plain, kt=plain, parts=[])
            new_entries.append(cur)
        elif cur is None: keep.append((k,t))
        else: cur['parts'].append((k,t))
    if new_entries:
        an['intro']=keep
        an['entries']=new_entries+an.get('entries',[])

def assemble(sections):
    tune_anatomy(sections)
    tune_overview(sections)
    restructure_timing(sections)
    restructure_skilltest(sections)
    split_cardtype_subs(sections)
    split_faq(sections)
    def S(k): return sections.get(k, dict(title=None,intro=[],entries=[]))
    if 'glossary' in sections:
        sections['glossary']['entries'] = sort_kr(sections['glossary']['entries'])
    opt = S('optional')['entries']
    env_names={'Current, Legacy, and Limited Environments','Current Environment','Legacy Environment','Limited Environment',
               'Regarding Parallels, Promos, Reward Cards, and Story Assets','Ultimatums and Boons'}
    env_entries=[e for e in opt if e['et'] in env_names]
    def find(name): return next((e for e in opt if e['et']==name), None)
    ult_hdr, boon_hdr, refr_hdr = find('Ultimatums'), find('Boons'), find('Refractions')
    def between(a,b):
        ia = opt.index(a)+1 if a in opt else 0
        ib = opt.index(b) if (b and b in opt) else len(opt)
        return opt[ia:ib]
    ults  = between(ult_hdr, boon_hdr) if ult_hdr else []
    boons = between(boon_hdr, refr_hdr) if boon_hdr else []
    refrs = between(refr_hdr, None) if refr_hdr else []

    rules_secs = [
        ('c2sec-overview', S('overview')), ('c2sec-glossary', S('glossary')),
        ('c2sec-fundamentals', S('fundamentals')), ('c2sec-timing_and_gameplay', S('timing')),
        ('c2sec-skill_test_timing', S('skilltest')), ('c2sec-initiation_sequence', S('initiation')),
        ('c2sec-scenario_setup', S('setup')), ('c2sec-card_anatomy', S('anatomy')),
        ('c2sec-campaign_play', S('campaign')), ('c2sec-deck_customization', S('deckcustom')),
        ('c2sec-icon_reference', S('iconref')), ('c2sec-encounter_set_icon_reference', S('encicon')),
        ('c2sec-quick_reference', S('quickref')),
    ]
    rules_secs=[(sid,s) for sid,s in rules_secs if s.get('title') or s.get('entries')]
    # 웹판 미수록: 조우 세트 아이콘 참조(XV)부터 끝까지 (한눈에 보는 요약 포함)
    DROP_RULES_SECS={'c2sec-encounter_set_icon_reference','c2sec-quick_reference'}
    rules_secs=[(sid,s) for sid,s in rules_secs if sid not in DROP_RULES_SECS]
    # 조우 세트 변형: 첫 엔트리를 섹션 헤더로 승격
    ev=S('encvar')
    if ev.get('entries') and not ev.get('title'):
        first=ev['entries'][0]
        ev=dict(title=first['kt'], intro=first['parts'], entries=ev['entries'][1:])
    ult_secs = [
        ('c2sec-optional_rules', dict(title=S('optional')['title'] or '선택 규칙', intro=S('optional')['intro'], entries=env_entries)),
        ('c2sec-ultimatums', dict(title='최후통첩', intro=(ult_hdr or {}).get('parts',[]), entries=ults)),
        ('c2sec-boons', dict(title='은총', intro=(boon_hdr or {}).get('parts',[]), entries=boons)),
        ('c2sec-refractions', dict(title='왜곡', intro=(refr_hdr or {}).get('parts',[]), entries=refrs)),
        ('c2sec-encounter_set_variation', ev),
    ]
    ult_secs=[(sid,s) for sid,s in ult_secs if s.get('title') or s.get('entries')]
    for secs in (rules_secs, ult_secs):
        for sid, sd in secs:
            sd['entries']=[e for e in sd.get('entries',[]) if e.get('et') not in DROP_ENTRIES]
    return {
      'rules': rules_secs,
      'errata': [('c2sec-notes_and_errata', S('errata'))],
      'faq': [('c2sec-frequently_asked_questions', S('faq'))],
      'ultimatums': ult_secs,
      'modreprints': [('c2sec-modified_reprints', S('modreprints'))],
    }

def panel_content_html(secs, header=None):
    set_title_map(secs)
    h = (header+'\n') if header else ''
    for sid, s in secs:
        for pe in s.get('pre', []):                   # 섹션 제목보다 앞서는 블록
            h += entry_html(pe)+'\n'
        h += section_html(sid, s.get('title') or '', s.get('intro',[]), s.get('entries',[]))+'\n'
    return h

def panel_toc_html(secs):
    return ''.join(toc_section(sid, s.get('title') or '', s.get('entries',[])) for sid,s in secs)

def replace_block(doc, open_tag, next_prefixes, inner):
    i = doc.find(open_tag)
    assert i>=0, '패널을 찾을 수 없음: '+open_tag
    j0 = i+len(open_tag)
    cands=[doc.find(p, j0) for p in next_prefixes]
    j = min(x for x in cands if x>=0)
    return doc[:i] + open_tag + '\n' + inner + '\n</div>\n' + doc[j:]

def find_latest_sheet():
    """스크립트 폴더/상위 폴더에서 '마도서 대조표' 시트를 가진 최신 *번역대조표*.xlsx 자동 탐색"""
    import os, glob
    here = os.path.dirname(os.path.abspath(__file__))
    dirs = [here, os.path.dirname(here), os.path.dirname(os.path.dirname(here)), os.getcwd()]
    cands = []
    for d in dict.fromkeys(dirs):
        for p in glob.glob(os.path.join(d, '*번역대조표*.xlsx')):
            b = os.path.basename(p)
            if b.startswith('~$') or b.startswith('.~'): continue
            cands.append(p)
    valid = []
    for p in sorted(set(cands), key=os.path.getmtime, reverse=True):
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            ok = '마도서 대조표' in wb.sheetnames
            wb.close()
            if ok: valid.append(p)
        except Exception:
            continue
    if not valid:
        sys.exit('오류: "마도서 대조표" 시트를 가진 번역대조표 xlsx를 찾지 못했습니다. 인자로 직접 지정해 주세요.')
    return valid[0]

def main():
    import os
    here = os.path.dirname(os.path.abspath(__file__))
    if len(sys.argv) >= 3:
        xlsx, target = sys.argv[1], sys.argv[2]
    elif len(sys.argv) == 2:
        # 인자 1개면: xlsx로 간주, 대상은 스크립트 옆 chapters.html
        xlsx, target = sys.argv[1], os.path.join(here, 'chapters.html')
    else:
        xlsx, target = find_latest_sheet(), os.path.join(here, 'chapters.html')
    print('시트  :', xlsx)
    print('대상  :', target)
    rows = parse(xlsx)
    sections, order = build_entries(rows)
    panels = assemble(sections)
    doc = open(target, encoding='utf-8').read()
    header = ('<div class="chapter-header"><h1>챕터 2 · 아컴 마도서</h1>'
              '<p>코리아보드게임즈와 공동 검수 중인 한국어 번역 초안 기준입니다. 검수 진행에 따라 내용이 변경될 수 있습니다.</p></div>')
    conf = [
      ('rules','<div class="content-panel grimoire-list" data-sub="rules" data-ch="c2">',header),
      ('errata','<div class="content-panel grimoire-list" data-sub="errata" data-ch="c2">',None),
      ('modreprints','<div class="content-panel grimoire-list" data-sub="modreprints" data-ch="c2">',None),
      ('faq','<div class="content-panel grimoire-list" data-sub="faq" data-ch="c2">',None),
      ('ultimatums','<div class="content-panel grimoire-list" data-sub="ultimatums" data-ch="c2">',None),
    ]
    for key, tag, hdr in conf:
        doc = replace_block(doc, tag, ['<div class="content-panel'], panel_content_html(panels[key], hdr))
    for key, tag in [
      ('rules','<div class="toc-panel toc" data-sub="rules" data-ch="c2">'),
      ('errata','<div class="toc-panel toc" data-sub="errata" data-ch="c2">'),
      ('modreprints','<div class="toc-panel toc" data-sub="modreprints" data-ch="c2">'),
      ('faq','<div class="toc-panel toc" data-sub="faq" data-ch="c2">'),
      ('ultimatums','<div class="toc-panel toc" data-sub="ultimatums" data-ch="c2">')]:
        doc = replace_block(doc, tag, ['<div class="toc-panel'], panel_toc_html(panels[key]))
    open(target,'w',encoding='utf-8').write(doc)
    n=sum(len(s['entries']) for s in sections.values())
    print('완료: 섹션 %d개, 엔트리 %d개 반영 -> %s' % (len(sections), n, target))

if __name__=='__main__':
    main()
